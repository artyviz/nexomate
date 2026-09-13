# excel/formatter.py
"""Apply professional formatting to exported Excel workbooks."""

from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# Color scheme
HEADER_FILL = PatternFill(start_color="111111", end_color="111111", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Calibri", size=10)
HIGH_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
MEDIUM_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
LOW_FILL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)


def apply_formatting(filepath: str | Path):
    """Apply professional formatting to an Excel workbook."""
    wb = load_workbook(filepath)

    for ws in wb.worksheets:
        # Format header row
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER

        # Format data rows
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=True)

            # Color-code by score level
            score_col = None
            for i, header_cell in enumerate(ws[1], 1):
                if header_cell.value == "Score Level":
                    score_col = i
                    break

            if score_col and len(row) >= score_col:
                level_cell = row[score_col - 1]
                level = (level_cell.value or "").upper()
                fill = None
                if level == "HIGH":
                    fill = HIGH_FILL
                elif level == "MEDIUM":
                    fill = MEDIUM_FILL
                elif level == "LOW":
                    fill = LOW_FILL

                if fill:
                    for cell in row:
                        cell.fill = fill

        # Auto-fit column widths (approximate)
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            adjusted_width = min(max_length + 4, 50)
            ws.column_dimensions[col_letter].width = adjusted_width

        # Freeze header row
        ws.freeze_panes = "A2"

    wb.save(filepath)
